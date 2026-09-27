"""The Word evidence report, laid out for a tender submission: cover with key figures, document control, contents,
executive summary, numbered sections with numbered figures and tables, landscape results appendix and a glossary.

Every figure in it is a model prediction and the wording says so. Built with python-docx; the few things python-docx
has no API for (fields, table borders and shading, repeating header rows) are written as WordprocessingML directly.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# Palette: navy text and table heads, one teal accent (the viewer's light-theme accent), cool neutrals.
NAVY, ACCENT, INK, MUTED, HAIR, BAND, CALLOUT = "1B2A41", "00707C", "1F2933", "5D6C7B", "D5DBE1", "F4F6F8", "EAF4F5"
FONT = "Calibri"
PORTRAIT_TEXT_CM = 21.0 - 2 * 2.2
LANDSCAPE_TEXT_CM = 29.7 - 2 * 2.0

VEHICLES = {"EXTERNAL_ROOFTOP_ANTENNA": "Passive rooftop antenna", "EDGE_RAIL_ACTIVE_ANTENNA": "EDGE Rail 5G active antenna",
            "PASSENGER_HANDSET_INSIDE_CARRIAGE": "Passenger handset inside the carriage"}
POLICIES = {"PACKET_BONDING": "Packet bonding (every usable link aggregated)", "WEIGHTED_LOAD_BALANCING": "Weighted load balancing",
            "FAILOVER": "Failover to the best single link", "CELLULAR_PRIMARY_STARLINK_BACKUP": "Cellular primary, satellite backup",
            "STARLINK_PRIMARY_CELLULAR_BACKUP": "Satellite primary, cellular backup"}
SERVICE_CLASSES = [("EXCELLENT", "80–100", "HD video streaming and video calls for every active passenger"),
                   ("GOOD", "60–79", "Video calls and streaming"),
                   ("USABLE", "40–59", "Web browsing, email and messaging"),
                   ("POOR", "15–39", "Intermittent; basic messaging only"),
                   ("OUTAGE", "0–14", "No usable service")]
CONFIDENCE_SCALE = [("0.90–1.00", "Directly measured or well validated"), ("0.70–0.89", "Strong source coverage and a calibrated model"),
                    ("0.40–0.69", "Prediction with partial infrastructure support"), ("0.10–0.39", "Sparse data or synthetic estimate"),
                    ("0.00", "Unknown")]
GLOSSARY = [
    ("Availability", "Share of the route over which a link is in service and good enough for the link manager to use."),
    ("Band", "One of five levels (Excellent to Very poor) used by the route heat maps; each metric's band edges are in its legend."),
    ("Combined onboard WAN", "The train's connection to the internet after the link manager has combined the usable mobile and satellite links."),
    ("Confidence", "0–1 score stating how much of an estimate rests on measured, predicted or synthetic inputs (section 7)."),
    ("DAS", "Distributed antenna system: in-tunnel mobile coverage. Tunnels are assumed to have none unless listed in the assumptions."),
    ("Handover", "The train's modem moving from one cell site to the next; briefly reduces throughput and raises latency."),
    ("Latency", "Round-trip delay of the combined onboard WAN, in milliseconds."),
    ("Link manager / policy", "The onboard router logic that decides which links carry traffic: bonding, load balancing or failover."),
    ("Median (P50)", "The value exceeded over half of the route."),
    ("Ofcom coverage prediction", "Mobile operators' predicted outdoor coverage by postcode, published through the Ofcom API."),
    ("OpenCellID", "Community-collected database of cell-site locations, used for serving-cell distance and handovers."),
    ("P10", "The value exceeded over 90 % of the route: a measure of the weak stretches rather than the average."),
    ("pp", "Percentage points: the difference between two percentages (96 % to 98 % is +2 pp)."),
    ("Packet bonding", "Aggregating several links at once so the train's capacity is close to their sum."),
    ("RSRP", "Reference Signal Received Power: the signal strength of an LTE/5G cell at the modem, in dBm. −80 dBm or more is excellent, below −110 dBm very poor."),
    ("Sample", "One point every 50 m along the railway at which every link and the passenger experience are estimated."),
    ("Service class", "Passenger Wi-Fi experience from per-user throughput, latency and loss: EXCELLENT, GOOD, USABLE, POOR or OUTAGE (section 1.3)."),
    ("Sky visibility", "Share of the sky above the satellite minimum elevation left open by terrain, cuttings, canopies and buildings."),
    ("Streaming-capable", "In the EXCELLENT or GOOD service class: video calls and streaming work."),
]


@dataclass
class Evidence:
    """Everything the document needs, assembled by report.build_report."""
    meta: dict
    route_name: str
    origin: str
    destination: str
    k: dict                                      # report.kpis()
    sec: pd.DataFrame                            # report.section_table()
    links: pd.DataFrame                          # report.link_table(), display names in "name"
    outages: pd.DataFrame                        # report.outage_stretches()
    shares: dict[str, list[float]]               # heat-map band shares (%) per metric
    charts: dict[str, bytes]
    assumptions: list[tuple[str, str]]
    sources: list[tuple[str, str, str]]
    validation: pd.DataFrame | None
    scenario_title: str                          # "Baseline configuration", a preset label or a train design
    scenario_detail: str                         # vehicle profile · policy (· weather)
    vehicle: str
    policy: str
    satcom: str
    passengers: str
    operators: list[str]
    design: dict | None
    claims: list[str]
    info: dict                                   # config/report.yaml
    baseline: dict | None = None                 # {"title", "k", "shares"} when this scenario is not the baseline
    generated: datetime = field(default_factory=lambda: datetime.now(timezone.utc))


# ------------------------------------------------------------------------------------------------ low-level helpers
def _el(tag: str, **attrs) -> OxmlElement:
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(f"w:{k}"), str(v))
    return e


# WordprocessingML fixes the order of property children; Word rejects a file that breaks it, so every raw insert goes
# through _put with its parent's sequence (only the elements this module writes, plus what python-docx writes).
_ORDER = {
    "pPr": ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr", "suppressLineNumbers", "pBdr", "shd",
            "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi",
            "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
            "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"],
    "rPr": ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline", "shadow", "emboss", "imprint",
            "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing", "w", "kern", "position", "sz", "szCs", "highlight", "u",
            "effect", "bdr", "shd", "fitText", "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath"],
    "tblPr": ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing",
              "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription"],
    "tcPr": ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark"],
    "trPr": ["cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter", "cantSplit", "trHeight", "tblHeader", "tblCellSpacing", "jc", "hidden"],
    "settings": ["updateFields", "hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids", "mathPr", "attachedSchema",
                 "themeFontLang", "clrSchemeMapping", "doNotIncludeSubdocsInStats", "doNotAutoCompressPictures", "forceUpgrade", "captions",
                 "readModeInkLockDown", "smartTagType", "schemaLibrary", "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol", "listSeparator"],
}


def _local(el) -> str:
    return el.tag.rsplit("}", 1)[-1]


def _put(parent, child, kind: str):
    """Insert child into parent at its schema position, replacing any existing element of the same name."""
    order, tag = _ORDER[kind], _local(child)
    for old in [c for c in parent if _local(c) == tag]:
        parent.remove(old)
    rank = order.index(tag)
    for existing in parent:
        et = _local(existing)
        if et in order and order.index(et) > rank:
            existing.addprevious(child)
            return child
    parent.append(child)
    return child


def _rgb(hex6: str) -> RGBColor:
    return RGBColor.from_string(hex6)


_UNIT = re.compile(r"(?<=[\d%)]) (?=(?:%|km|Mbps|ms|dBm|dB|min|h|m|points)\b)")


def _nb(text: str) -> str:
    """Keep numbers with their units on one line (a non-breaking space before %, km, Mbps, ms, dB ...)."""
    return _UNIT.sub("\u00a0", text)


def _and(items: list[str]) -> str:
    items = [i for i in items if i]
    return items[0] if len(items) == 1 else (", ".join(items[:-1]) + " and " + items[-1] if items else "")


def _a(noun: str) -> str:
    return ("an " if noun[:1].lower() in "aeiou" else "a ") + noun


def _run(par, text: str, size: float | None = None, bold: bool = False, color: str | None = None, italic: bool = False, caps: bool = False):
    r = par.add_run(_nb(text))
    if size:
        r.font.size = Pt(size)
    r.font.bold = bold or None
    r.font.italic = italic or None
    if color:
        r.font.color.rgb = _rgb(color)
    if caps:
        r.font.all_caps = True
    return r


def _field(par, instr: str, cached: str = "", size: float | None = None, bold: bool = False, color: str | None = None):
    """A Word field (PAGE, NUMPAGES, SEQ ...) with a cached result, so it reads correctly before Word updates it."""
    for kind, text in (("begin", None), ("instr", instr), ("separate", None), ("text", cached), ("end", None)):
        r = par.add_run()
        if size:
            r.font.size = Pt(size)
        r.font.bold = bold or None
        if color:
            r.font.color.rgb = _rgb(color)
        if kind == "instr":
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {text} "; r._r.append(it)
        elif kind == "text":
            r.text = text
        else:
            r._r.append(_el("w:fldChar", fldCharType=kind))


def _border(par, side: str, color: str, size: int = 4, space: int = 4) -> None:
    pPr = par._p.get_or_add_pPr()
    bdr = pPr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = _put(pPr, OxmlElement("w:pBdr"), "pPr")
    bdr.append(_el(f"w:{side}", val="single", sz=size, space=space, color=color))


def _shade(cell, fill: str) -> None:
    _put(cell._tc.get_or_add_tcPr(), _el("w:shd", val="clear", color="auto", fill=fill), "tcPr")


def _cell_borders(cell, **sides) -> None:
    b = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):             # schema order
        if side in sides:
            b.append(_el(f"w:{side}", **sides[side]))
    _put(cell._tc.get_or_add_tcPr(), b, "tcPr")


def _fixed(table, total_cm: float) -> None:
    """Fixed layout at an exact width, so Word and LibreOffice keep the column widths instead of autofitting."""
    table.autofit = False
    tblPr = table._tbl.tblPr
    _put(tblPr, _el("w:tblW", w=int(total_cm * 567), type="dxa"), "tblPr")
    _put(tblPr, _el("w:tblLayout", type="fixed"), "tblPr")


def _fmt(v, spec: str = "") -> str:
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return "–"
    if spec and isinstance(v, (int, float, np.integer, np.floating)):
        return format(float(v), spec).replace("-", "−")
    return str(v)


def _share(v: float) -> str:
    """A share of the route: whole percent, except that nearly-all and nearly-none keep a decimal (99.6 % is not 100 %)."""
    return f"{v:.1f}" if (99.5 <= v < 100 or 0 < v < 0.5) else f"{v:.0f}"


def _hours(minutes: float) -> str:
    h, m = divmod(int(round(minutes)), 60)
    return f"{h} h {m:02d} min" if h else f"{m} min"


# ------------------------------------------------------------------------------------------------ the document
class _Report:
    def __init__(self, ev: Evidence):
        self.ev = ev
        self.doc = Document()
        self.figures = 0
        self.tables = 0
        self.sections = 0
        self.subsections = 0
        self.toc: list[tuple[int, str]] = []
        self._styles()

    # ---- styles and page furniture ---------------------------------------------------------------------------
    def _styles(self) -> None:
        st = self.doc.styles
        normal = st["Normal"]
        normal.font.name = FONT; normal.font.size = Pt(10); normal.font.color.rgb = _rgb(INK)
        rpr = normal.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts"); rpr.append(fonts)
        for a in ("ascii", "hAnsi", "cs", "eastAsia"):
            fonts.set(qn(f"w:{a}"), FONT)
        _put(rpr, _el("w:lang", val="en-GB"), "rPr")
        pf = normal.paragraph_format
        pf.space_after = Pt(6); pf.line_spacing = 1.12
        for name, size, before, after in (("Heading 1", 16, 18, 8), ("Heading 2", 12, 12, 4), ("Heading 3", 10.5, 10, 3)):
            h = st[name]
            h.font.name = FONT; h.font.size = Pt(size); h.font.bold = True; h.font.italic = False; h.font.color.rgb = _rgb(NAVY)
            hr = h.element.get_or_add_rPr()
            hf = hr.find(qn("w:rFonts"))
            if hf is None:
                hf = OxmlElement("w:rFonts"); hr.append(hf)
            for a in ("ascii", "hAnsi", "cs", "eastAsia"):
                hf.set(qn(f"w:{a}"), FONT)
            for a in ("asciiTheme", "hAnsiTheme", "cstheme", "eastAsiaTheme"):
                hf.attrib.pop(qn(f"w:{a}"), None)
            h.paragraph_format.space_before = Pt(before); h.paragraph_format.space_after = Pt(after)
            h.paragraph_format.keep_with_next = True
        cap = st["Caption"]
        cap.font.name = FONT; cap.font.size = Pt(8.5); cap.font.bold = False; cap.font.italic = False; cap.font.color.rgb = _rgb(MUTED)
        cap.paragraph_format.space_before = Pt(3); cap.paragraph_format.space_after = Pt(12)
        for name, indent, bold in (("toc 1", 0, True), ("toc 2", 0.9, False)):
            s = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            s.base_style = normal
            s.font.bold = bold
            s.paragraph_format.left_indent = Cm(indent)
            s.paragraph_format.space_before = Pt(6 if bold else 0); s.paragraph_format.space_after = Pt(2)
            s.paragraph_format.tab_stops.add_tab_stop(Cm(PORTRAIT_TEXT_CM), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        _put(self.doc.settings.element, _el("w:updateFields", val="true"), "settings")   # Word refreshes the contents page on opening
        zoom = self.doc.settings.element.find(qn("w:zoom"))
        if zoom is not None and zoom.get(qn("w:percent")) is None:
            zoom.set(qn("w:percent"), "100")                  # required by the schema; python-docx's template omits it
        cp = self.doc.core_properties
        cp.title = f"Onboard connectivity performance: {self.ev.route_name}"
        cp.subject = f"{self.ev.scenario_title}. Model prediction."
        cp.author = cp.last_modified_by = self.ev.info.get("prepared_by") or "Train Link Simulator"
        cp.keywords = "tender evidence; onboard connectivity; model prediction"
        cp.comments = ""
        cp.category = self.ev.info.get("classification") or ""

    def _page(self, section, landscape: bool) -> None:
        section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = (Cm(29.7), Cm(21.0)) if landscape else (Cm(21.0), Cm(29.7))
        side = Cm(2.0 if landscape else 2.2)
        section.left_margin = section.right_margin = side
        section.top_margin, section.bottom_margin = Cm(2.2), Cm(2.0)
        section.header_distance, section.footer_distance = Cm(1.0), Cm(1.0)
        width = LANDSCAPE_TEXT_CM if landscape else PORTRAIT_TEXT_CM
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False
        ev, info = self.ev, self.ev.info
        section.different_first_page_header_footer = section is self.doc.sections[0]   # only the cover goes without
        hp = section.header.paragraphs[0]
        hp.text = ""
        self._tabs(hp, width)
        _run(hp, f"Onboard connectivity performance  ·  {ev.route_name}", 8, color=MUTED)
        if info.get("classification"):
            hp.add_run("\t")
            _run(hp, info["classification"], 8, bold=True, color=NAVY, caps=True)
        _border(hp, "bottom", HAIR, 4, 4)
        fp = section.footer.paragraphs[0]
        fp.text = ""
        self._tabs(fp, width)
        _border(fp, "top", HAIR, 4, 4)
        _run(fp, self.reference, 8, color=MUTED)
        _run(fp, f"   ·   Version {info.get('version') or '1.0'}   ·   Model prediction", 8, color=MUTED)
        fp.add_run("\t")
        _run(fp, "Page ", 8, color=MUTED)
        _field(fp, "PAGE", "1", 8, color=MUTED)
        _run(fp, " of ", 8, color=MUTED)
        _field(fp, "NUMPAGES", "1", 8, color=MUTED)

    @staticmethod
    def _tabs(par, width_cm: float) -> None:
        """One right-aligned stop at the text edge; clears the Header/Footer styles' centre (8.25 cm) and right (16.51 cm) stops."""
        ts = par.paragraph_format.tab_stops
        for pos in (Cm(8.255), Cm(16.51)):
            ts.add_tab_stop(pos, WD_TAB_ALIGNMENT.CLEAR)
        ts.add_tab_stop(Cm(width_cm), WD_TAB_ALIGNMENT.RIGHT)

    @property
    def reference(self) -> str:
        preset = self.ev.meta.get("scenario_id", "baseline")
        return f"TLS-{self.ev.meta['route']['id'].upper()}-{preset.upper()}-{self.ev.generated.strftime('%Y%m%d')}"

    # ---- blocks ----------------------------------------------------------------------------------------------
    def para(self, text: str = "", size: float | None = None, color: str | None = None, bold: bool = False, after: float | None = None,
             align=None, keep: bool = False, italic: bool = False):
        p = self.doc.add_paragraph()
        if text:
            _run(p, text, size, bold, color, italic)
        if after is not None:
            p.paragraph_format.space_after = Pt(after)
        if align is not None:
            p.alignment = align
        if keep:
            p.paragraph_format.keep_with_next = True
        return p

    def rich(self, parts: list[tuple[str, bool]], after: float | None = None):
        p = self.doc.add_paragraph()
        for text, bold in parts:
            _run(p, text, bold=bold)
        if after is not None:
            p.paragraph_format.space_after = Pt(after)
        return p

    def bullets(self, items: list[str | list[tuple[str, bool]]]) -> None:
        for it in items:
            p = self.doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            for text, bold in ([(it, False)] if isinstance(it, str) else it):
                _run(p, text, bold=bold)

    def page_break(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def front_heading(self, text: str) -> None:
        """A heading that is not numbered and not in the contents (document control, contents)."""
        p = self.para(text, 16, NAVY, True, after=10, keep=True)
        p.paragraph_format.space_before = Pt(0)

    def h1(self, text: str, number: str | None = None, page_break: bool = False):
        if number is None:
            self.sections += 1
            self.subsections = 0
            number = str(self.sections)
        label = f"{number}   {text}" if number else text
        h = self.doc.add_heading(label, 1)
        if page_break:
            h.paragraph_format.page_break_before = True
        self.toc.append((1, label))
        return h

    def h2(self, text: str):
        self.subsections += 1
        label = f"{self.sections}.{self.subsections}   {text}"
        self.doc.add_heading(label, 2)
        self.toc.append((2, label))

    def caption(self, kind: str, text: str, before_content: bool = False) -> None:
        if kind == "Figure":
            self.figures += 1; n = self.figures
        else:
            self.tables += 1; n = self.tables
        p = self.doc.add_paragraph(style="Caption")
        _run(p, f"{kind} ", bold=True, color=NAVY)
        _field(p, f"SEQ {kind} \\* ARABIC", str(n), bold=True, color=NAVY)
        _run(p, f"   {text}")
        if before_content:
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)

    def figure(self, png: bytes, caption: str, width_cm: float = PORTRAIT_TEXT_CM) -> None:
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(0)
        p.add_run().add_picture(io.BytesIO(png), width=Cm(width_cm))
        self.caption("Figure", caption)

    def table(self, columns: list[tuple[str, float, str]], rows: list[list[str]], caption: str | None = None, size: float = 8.5,
              bold_first: bool = False, header: bool = True, total_cm: float = PORTRAIT_TEXT_CM):
        """columns: (header, width in cm, 'l' | 'r' | 'c'). Header row repeats on every page; rows never split."""
        if caption:
            self.caption("Table", caption, before_content=True)
        widths = [c[1] for c in columns]
        scale = total_cm / sum(widths)
        widths = [w * scale for w in widths]
        t = self.doc.add_table(rows=(1 if header else 0) + len(rows), cols=len(columns))
        _fixed(t, total_cm)
        tblPr = t._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            if side in ("left", "right", "insideV"):
                borders.append(_el(f"w:{side}", val="nil"))
            else:
                borders.append(_el(f"w:{side}", val="single", sz=4, space=0, color=HAIR))
        _put(tblPr, borders, "tblPr")
        mar = OxmlElement("w:tblCellMar")
        for side, v in (("top", 40), ("left", 80), ("bottom", 40), ("right", 80)):
            mar.append(_el(f"w:{side}", w=v, type="dxa"))
        _put(tblPr, mar, "tblPr")
        for j, w in enumerate(widths):
            t.columns[j].width = Cm(w)                      # the grid: LibreOffice and Word lay out from it
        align = {"l": WD_ALIGN_PARAGRAPH.LEFT, "r": WD_ALIGN_PARAGRAPH.RIGHT, "c": WD_ALIGN_PARAGRAPH.CENTER}
        all_rows = ([[c[0] for c in columns]] if header else []) + rows
        for i, (row, values) in enumerate(zip(t.rows, all_rows)):
            trPr = row._tr.get_or_add_trPr()
            _put(trPr, _el("w:cantSplit", val="true"), "trPr")
            is_head = header and i == 0
            if is_head:
                _put(trPr, _el("w:tblHeader", val="true"), "trPr")
            for j, (cell, value) in enumerate(zip(row.cells, values)):
                cell.width = Cm(widths[j])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
                p.alignment = align[columns[j][2]]
                if is_head:
                    _shade(cell, NAVY)
                    _run(p, str(value), size, True, "FFFFFF")
                else:
                    if (i - (1 if header else 0)) % 2 == 1:
                        _shade(cell, BAND)
                    _run(p, str(value), size, bold_first and j == 0, INK)
                if len(all_rows) <= 16 and i < len(all_rows) - 1:
                    p.paragraph_format.keep_with_next = True   # a short table stays on one page
        self.doc.add_paragraph().paragraph_format.space_after = Pt(4)
        return t

    def kv(self, pairs: list[tuple[str, str]], caption: str | None = None, label_cm: float = 5.2, size: float = 9) -> None:
        self.table([("", label_cm, "l"), ("", PORTRAIT_TEXT_CM - label_cm, "l")], [[a, b] for a, b in pairs if b], caption, size, bold_first=True, header=False)

    def callout(self, title: str, text: str) -> None:
        t = self.doc.add_table(rows=1, cols=1)
        _fixed(t, PORTRAIT_TEXT_CM)
        t.columns[0].width = Cm(PORTRAIT_TEXT_CM)
        cell = t.rows[0].cells[0]
        cell.width = Cm(PORTRAIT_TEXT_CM)
        _shade(cell, CALLOUT)
        _cell_borders(cell, left={"val": "single", "sz": 24, "space": 0, "color": ACCENT}, top={"val": "nil"}, bottom={"val": "nil"}, right={"val": "nil"})
        mar = OxmlElement("w:tcMar")
        for side, v in (("top", 120), ("left", 200), ("bottom", 120), ("right", 200)):
            mar.append(_el(f"w:{side}", w=v, type="dxa"))
        _put(cell._tc.get_or_add_tcPr(), mar, "tcPr")
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(3)
        _run(p, title, 10, True, NAVY)
        q = cell.add_paragraph()
        q.paragraph_format.space_after = Pt(0)
        _run(q, text, 9.5, color=INK)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # ---- the contents page -----------------------------------------------------------------------------------
    def contents(self, anchor) -> None:
        """Fill the contents placeholder: a TOC field whose cached entries list every heading, so it reads correctly
        before Word refreshes it (Word adds the page numbers when it updates fields on opening)."""
        entries = self.toc or [(1, "")]
        paras = []
        for level, text in entries:
            p = anchor.insert_paragraph_before(style=f"toc {level}")
            paras.append((p, text))
        first = paras[0][0]
        for kind, text in (("begin", None), ("instr", 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None)):
            r = first.add_run()
            if kind == "instr":
                it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {text} "; r._r.append(it)
            else:
                r._r.append(_el("w:fldChar", fldCharType=kind))
        for p, text in paras:
            p.add_run(text)
        paras[-1][0].add_run()._r.append(_el("w:fldChar", fldCharType="end"))
        note = anchor
        note.paragraph_format.space_before = Pt(14)
        _run(note, "Page numbers are filled in when Word updates the document's fields (it offers to on opening; or press F9 in the contents).", 8, color=MUTED, italic=True)

    # ---- pages -----------------------------------------------------------------------------------------------
    def cover(self) -> None:
        ev, k, info = self.ev, self.ev.k, self.ev.info
        self.para(after=70)
        p = self.para("Tender evidence  ·  onboard connectivity", 9.5, ACCENT, True, after=10)
        p.runs[0].font.all_caps = True
        self.para("Onboard connectivity performance", 28, NAVY, True, after=6).paragraph_format.line_spacing = 1.0
        self.para(ev.route_name, 17, INK, after=2)
        self.para(f"{ev.origin} to {ev.destination}   ·   {k['route_length_km']:.0f} km   ·   {_hours(k['journey_minutes'])}", 11, MUTED, after=18)
        p = self.para(after=2)
        _run(p, "Scenario  ", 9, True, MUTED, caps=True)
        _run(p, ev.scenario_title, 12, True, NAVY)
        self.para(ev.scenario_detail, 9.5, MUTED, after=28)
        tiles = [(_share(k["streaming_share_pct"]), "%", "of the journey supports\nvideo calls and streaming"),
                 (_share(k["usable_share_pct"]), "%", "usable or better\n(browsing, email, messaging)"),
                 (f"{k['bonded_median_mbps']:.0f}", "Mbps", "median combined\nonboard throughput"),
                 (f"{k['outage_km']:.1f}", "km", f"predicted loss of service\n({k['outage_share_pct']:.1f} % of the route)")]
        t = self.doc.add_table(rows=1, cols=4)
        _fixed(t, PORTRAIT_TEXT_CM)
        for j in range(4):
            t.columns[j].width = Cm(PORTRAIT_TEXT_CM / 4)
        for cell, (big, unit, small) in zip(t.rows[0].cells, tiles):
            cell.width = Cm(PORTRAIT_TEXT_CM / 4)
            _cell_borders(cell, top={"val": "single", "sz": 18, "space": 0, "color": ACCENT}, left={"val": "nil"}, bottom={"val": "nil"}, right={"val": "nil"})
            a = cell.paragraphs[0]
            a.paragraph_format.space_before = Pt(6); a.paragraph_format.space_after = Pt(2)
            _run(a, big, 24, True, NAVY)
            _run(a, "\u00a0" + unit, 12, True, NAVY)
            for line in small.split("\n"):
                b = cell.add_paragraph(); b.paragraph_format.space_after = Pt(0); b.paragraph_format.line_spacing = 1.0
                _run(b, line, 8.5, color=MUTED)
        rows = [("Prepared for", info.get("prepared_for", "")), ("Prepared by", info.get("prepared_by", "")),
                ("Tender reference", info.get("tender_reference", "")), ("Document reference", self.reference),
                ("Version", info.get("version") or "1.0"), ("Date of issue", ev.generated.strftime("%d %B %Y").lstrip("0")),
                ("Status", "Model prediction" + ("; validated against field measurements" if ev.validation is not None and len(ev.validation) else "; not yet validated against field measurements")),
                ("Classification", info.get("classification", ""))]
        rows = [(a, b) for a, b in rows if b]
        self.para(after=max(40, 330 - 17 * len(rows)))            # the details sit towards the foot of the cover
        self.kv(rows, label_cm=4.2, size=9)

    def document_control(self) -> None:
        ev, info = self.ev, self.ev.info
        self.page_break()
        self.front_heading("Document control")
        self.kv([("Title", f"Onboard connectivity performance: {ev.route_name}"), ("Scenario", f"{ev.scenario_title} ({ev.scenario_detail})"),
                 ("Document reference", self.reference), ("Version", info.get("version") or "1.0"),
                 ("Date of issue", ev.generated.strftime("%d %B %Y").lstrip("0")), ("Prepared by", info.get("prepared_by", "")),
                 ("Prepared for", info.get("prepared_for", "")), ("Tender reference", info.get("tender_reference", "")),
                 ("Contact", info.get("contact", "")), ("Classification", info.get("classification", "")),
                 ("Produced with", f"Train Link Simulator, model version {ev.meta.get('model_version', '')}"
                                   + (f"; route data built {ev.meta['built']}" if ev.meta.get("built") else ""))], "Document details")
        self.table([("Version", 2.0, "l"), ("Date", 3.2, "l"), ("Description", 11.4, "l")],
                   [[info.get("version") or "1.0", ev.generated.strftime("%d %b %Y").lstrip("0"),
                     f"Issued for tender: {ev.scenario_title.lower() if ev.scenario_title.startswith('Baseline') else ev.scenario_title} on the {ev.route_name}"]],
                   "Revision history")
        self.callout("Basis of preparation",
                     "The figures in this document are predictions from a simulation of the route, not measurements of a deployed system. "
                     "They combine the railway's geometry and terrain, the mobile operators' published coverage predictions, cell-site data, a "
                     "satellite visibility model and the onboard configuration described in section 5. Each estimate carries a confidence value "
                     "stating how much of it rests on measured, predicted or synthetic inputs (section 7), and section 8 states what has been "
                     "checked against field measurements. Assumptions and limitations are listed in section 9.")

    def contents_page(self):
        self.page_break()
        self.front_heading("Contents")
        return self.doc.add_paragraph()

    def executive_summary(self) -> None:
        ev, k, sh = self.ev, self.ev.k, self.ev.shares
        self.h1("Executive summary", number="", page_break=True)
        whom = f" for {ev.info['prepared_for']}" if ev.info.get("prepared_for") else ""
        ref = f" ({ev.info['tender_reference']})" if ev.info.get("tender_reference") else ""
        sat = f"the {ev.satcom} satellite link" if ev.satcom != "no satellite link" else "no satellite link"
        self.para(f"This document sets out the predicted onboard connectivity{whom}{ref} for passengers travelling on the {ev.route_name} "
                  f"between {ev.origin} and {ev.destination}: {k['route_length_km']:.0f} km in {_hours(k['journey_minutes'])}. It models "
                  f"{'the baseline configuration' if ev.scenario_title.startswith('Baseline') else 'the ' + ev.scenario_title + ' scenario'}, "
                  f"{_a(ev.vehicle.lower() if not ev.vehicle.startswith('EDGE') else ev.vehicle)} with {ev.policy.split(' (')[0].lower()}, "
                  f"across the {_and(ev.operators)} mobile networks and {sat}.")
        self.doc.add_heading("Key findings", 3)
        findings: list[list[tuple[str, bool]]] = [
            [("Video calls and streaming ", False), (f"are supported over {_share(k['streaming_share_pct'])} % of the journey", True),
             (f", and the service is usable or better over {_share(k['usable_share_pct'])} %.", False)],
            [("The combined onboard connection delivers ", False), (f"a median of {k['bonded_median_mbps']:.0f} Mbps", True),
             (f"; 90 % of the route receives at least {k['bonded_p10_mbps']:.0f} Mbps, and latency stays below 60 ms over "
              f"{sh['latency'][0] + sh['latency'][1]:.0f} % of it.", False)],
        ]
        if k["outage_km"] > 0 and len(ev.outages):
            o = ev.outages.iloc[0]
            where = f" in {o['tunnel']}" if o["tunnel"] else ""
            findings.append([("Predicted loss of service ", False), (f"totals {k['outage_km']:.1f} km ({k['outage_share_pct']:.1f} % of the route)", True),
                             (f"; the longest continuous gap is {o['length_km']:.1f} km, between {o['from']} and {o['to']}{where}.", False)])
        else:
            findings.append([("No loss of service is predicted", True), (" anywhere on the route.", False)])
        weak = ev.sec.sort_values(["streaming_share_pct", "outage_km"], ascending=[True, False]).head(3)
        findings.append([("The strongest mobile network is ", False), (f"Good or better (≥ −90 dBm) over {sh['signal'][0] + sh['signal'][1]:.0f} % of the route", True),
                         ("; the weakest sections for passengers are " + _and([f"{r['from']} – {r['to']}" for _, r in weak.iterrows()]) + ".", False)])
        if ev.baseline:
            b = ev.baseline["k"]

            def moves(what: str, before: float, after: float, unit: str, fmt: str) -> str:
                a, z = format(before, fmt), format(after, fmt)
                return f"{what} is unchanged at {z}{unit}" if a == z else f"{what} {'rises' if after > before else 'falls'} from {a}{unit} to {z}{unit}"
            findings.append([(f"Compared with the baseline ({ev.baseline['title']}), ", False),
                             (moves("the streaming-capable share of the journey", b["streaming_share_pct"], k["streaming_share_pct"], " %", ".1f"), True),
                             (", " + moves("median throughput", b["bonded_median_mbps"], k["bonded_median_mbps"], " Mbps", ".0f") + " and "
                              + moves("predicted loss of service", b["outage_km"], k["outage_km"], " km", ".1f") + ".", False)])
        validated = ev.validation is not None and len(ev.validation)
        findings.append([("The estimates carry ", False), (f"a mean confidence of {k['mean_confidence']:.2f}", True),
                         (" on a 0–1 scale" + ("; they have been compared with field measurements (section 8)." if validated else
                                               "; they have not yet been validated against field measurements (section 8)."), False)])
        self.bullets(findings)
        if ev.baseline:
            b, bs = ev.baseline["k"], ev.baseline["shares"]

            def row(label, before, after, unit, fmt=".0f", better_high=True):
                change_unit = " pp" if unit == " %" else unit           # a change in a share is in percentage points
                d = float(format(after, fmt)) - float(format(before, fmt))
                verdict = "" if d == 0 else ("better" if (d > 0) == better_high else "worse")
                change = "no change" if d == 0 else f"{'+' if d > 0 else '−'}{format(abs(d), fmt)}{change_unit}  ({verdict})"
                return [label, f"{format(before, fmt)}{unit}", f"{format(after, fmt)}{unit}", change]
            self.table([("Metric", 6.4, "l"), ("Baseline", 2.9, "r"), ("This scenario", 3.0, "r"), ("Change", 4.3, "r")], [
                row("Streaming-capable share of the journey", b["streaming_share_pct"], k["streaming_share_pct"], " %", ".1f"),
                row("Usable-or-better share of the journey", b["usable_share_pct"], k["usable_share_pct"], " %", ".1f"),
                row("Predicted loss of service", b["outage_km"], k["outage_km"], " km", ".1f", False),
                row("Median combined throughput", b["bonded_median_mbps"], k["bonded_median_mbps"], " Mbps"),
                row("Throughput exceeded over 90 % of the route (P10)", b["bonded_p10_mbps"], k["bonded_p10_mbps"], " Mbps"),
                row("Mean latency", b["latency_mean_ms"], k["latency_mean_ms"], " ms", ".0f", False),
                row("Journey where the strongest network is Good or better (≥ −90 dBm)", bs["signal"][0] + bs["signal"][1], sh["signal"][0] + sh["signal"][1], " %", ".1f"),
            ], f"This scenario compared with the baseline ({ev.baseline['title']})")
        self.table([("Section", 6.2, "l"), ("Length", 1.9, "r"), ("Streaming-capable", 2.8, "r"), ("Loss of service", 2.4, "r"), ("Least available link", 3.3, "l")],
                   [[f"{r['from']} – {r['to']}", f"{r['length_km']:.1f} km", f"{r['streaming_share_pct']:.0f} %", f"{r['outage_km']:.1f} km", r["weakest_name"]]
                    for _, r in weak.iterrows()], "Sections with the weakest predicted passenger experience")

    def introduction(self) -> None:
        ev, k = self.ev, self.ev.k
        self.h1("Introduction", page_break=True)
        self.h2("Purpose")
        self.para("This document provides evidence of the onboard connectivity passengers can expect on the route and service below, for use in "
                  "assessing the proposed onboard system. It reports where along the journey video calls, streaming and browsing will work, the "
                  "throughput and latency of the train's connection, how each mobile network and the satellite link contribute, and how much "
                  "confidence each figure deserves.")
        self.h2("Route, service and system modelled")
        calling = [s["name"] for s in ev.meta["stations"] if s.get("stop", True)]
        self.kv([("Route", ev.route_name), ("Between", f"{ev.origin} and {ev.destination}"),
                 ("Length and journey time", f"{k['route_length_km']:.1f} km; {_hours(k['journey_minutes'])} from departure to arrival"),
                 ("Calling pattern", ", ".join(calling)),
                 ("Scenario", ev.scenario_title), ("Onboard antenna", ev.vehicle), ("Link management", ev.policy),
                 ("Mobile networks", ", ".join(ev.operators)), ("Satellite", ev.satcom), ("Passenger load", ev.passengers),
                 ("Resolution", f"Every {ev.meta['route']['sample_spacing_m']} m along the railway ({ev.meta['n_samples']:,} points)")],
                "Scope of the assessment")
        self.h2("How to read this document")
        self.para("Passenger experience is reported in five service classes, derived from the throughput each active passenger receives, "
                  "latency and packet loss:", keep=True)
        self.table([("Class", 2.6, "l"), ("Score", 1.8, "c"), ("What passengers can do", 12.2, "l")], [list(r) for r in SERVICE_CLASSES],
                   "Passenger Wi-Fi service classes", bold_first=True)
        self.bullets(["Shares of the journey are by distance. “Streaming-capable” means the EXCELLENT or GOOD class.",
                      "P10 is the value exceeded over 90 % of the route, so it describes the weak stretches rather than the average.",
                      "Route heat maps use five bands per metric, darker meaning worse; each legend states the band edges and the share of the route in each.",
                      "Every figure is a prediction with a confidence value (section 7); terms are defined in the glossary (Appendix B)."])

    def performance(self) -> None:
        ev, k, sh = self.ev, self.ev.k, self.ev.shares
        self.h1("Predicted performance along the route", page_break=True)
        self.h2("Headline results")
        self.table([("Measure", 11.2, "l"), ("Prediction", 5.4, "r")], [
            ["Journey supporting video calls and streaming (EXCELLENT or GOOD)", f"{k['streaming_share_pct']:.1f} %"],
            ["Journey usable or better", f"{k['usable_share_pct']:.1f} %"],
            ["Predicted loss of service (OUTAGE class)", f"{k['outage_km']:.1f} km  ({k['outage_share_pct']:.1f} %)"],
            ["Combined onboard throughput: mean / median", f"{k['bonded_mean_mbps']:.0f} / {k['bonded_median_mbps']:.0f} Mbps"],
            ["Combined onboard throughput exceeded over 90 % of the route (P10)", f"{k['bonded_p10_mbps']:.0f} Mbps"],
            ["Throughput per active passenger (mean)", f"{k['per_user_mean_mbps']:.2f} Mbps"],
            ["Latency (mean, when connected)", f"{k['latency_mean_ms']:.0f} ms"],
            ["Journey with latency below 60 ms", f"{sh['latency'][0] + sh['latency'][1]:.0f} %"],
            ["Journey where the strongest network is Good or better (≥ −90 dBm)", f"{sh['signal'][0] + sh['signal'][1]:.0f} %"],
            ["Passenger Wi-Fi service score (mean, 0–100)", f"{k['wifi_score_mean']:.0f}"],
            ["Mean confidence of the estimates (0–1)", f"{k['mean_confidence']:.2f}"],
        ], "Headline predictions for the whole journey")
        self.h2("Throughput along the route")
        self.para("Figure 1 shows the combined onboard throughput from origin to destination: the line is the median over a short moving "
                  "window (stated in the legend) and the shaded band the range between the 10th and 90th percentiles within it. Calling "
                  "points are marked along the top; tunnels are shaded.", keep=True)
        self.figure(ev.charts["capacity"], "Predicted combined onboard throughput along the route")
        self.h2("Passenger service by section")
        self.figure(ev.charts["sections"], "Predicted passenger Wi-Fi service class by station-to-station section (share of each section's length)")
        self.h2("Route heat maps")
        self.para("Figure 3 maps where connectivity is strong and where it is weak, darker meaning worse. Signal is the strongest mobile network "
                  "received at the train with this scenario's antenna; throughput and latency are for the combined onboard connection, which "
                  "also draws on the satellite link. At page scale each short stretch of line shows the band that at least half of it reaches; "
                  f"the legend shares count every {ev.meta['route']['sample_spacing_m']} m point. Figure 4 breaks the signal down by network.", keep=True)
        self.figure(ev.charts["heatmap"], "Route heat maps: predicted signal strength, throughput and latency (darker = worse)")
        self.figure(ev.charts["networks"], "Predicted signal strength of each mobile network along the route (RSRP at the train, darker = weaker); "
                                           "the bottom row is the strongest of them, as mapped in Figure 3")

    def sections_and_links(self) -> None:
        ev = self.ev
        self.h1("Station-to-station performance")
        self.para("Each row is the section between consecutive calling points. Throughput and latency are for the combined onboard connection; "
                  "streaming-capable is the share of the section's length in the EXCELLENT or GOOD class; the least available link is the "
                  "network or satellite link usable over the smallest share of the section. Appendix A adds per-passenger throughput, the "
                  "usable share, tunnels and confidence for each section.")
        self.table([("Section", 5.6, "l"), ("km", 1.2, "r"), ("min", 1.2, "r"), ("Mean Mbps", 1.6, "r"), ("P10 Mbps", 1.5, "r"),
                    ("Latency ms", 1.6, "r"), ("Streaming %", 1.8, "r"), ("Loss km", 1.4, "r"), ("Least available", 2.1, "l")],
                   [[f"{r['from']} – {r['to']}", _fmt(r["length_km"], ".1f"), _fmt(r["minutes"], ".0f"), _fmt(r["bonded_mean_mbps"], ".0f"),
                     _fmt(r["bonded_p10_mbps"], ".0f"), _fmt(r["latency_mean_ms"], ".0f"), _fmt(r["streaming_share_pct"], ".0f"),
                     _fmt(r["outage_km"], ".1f"), r["weakest_name"]] for _, r in ev.sec.iterrows()],
                   "Predicted performance by station-to-station section", size=8)
        self.h1("Individual links")
        self.para("How each mobile network and the satellite link performs on its own, before the link manager combines them. Availability "
                  "is the share of the route over which the link is usable; capacities are for the times it is available.", keep=True)
        self.figure(ev.charts["links"], "Availability and median capacity of each link")
        self.table([("Link", 2.6, "l"), ("Type", 2.0, "l"), ("Available", 1.8, "r"), ("Median Mbps", 1.9, "r"), ("P10 Mbps", 1.6, "r"),
                    ("Latency ms", 1.8, "r"), ("Handovers", 1.8, "r"), ("Unavailable km", 2.0, "r"), ("Confidence", 1.8, "r")],
                   [[r["name"], "Mobile" if r["type"] == "cellular" else "Satellite", f"{r['availability_pct']:.0f} %", _fmt(r["p50_capacity_mbps"], ".0f"),
                     _fmt(r["p10_capacity_mbps"], ".0f"), _fmt(r["mean_latency_ms"], ".0f"), _fmt(r["handover_events"]), _fmt(r["unavailable_km"], ".1f"),
                     _fmt(r["mean_confidence"], ".2f")] for _, r in ev.links.iterrows()],
                   "Predicted performance of each link", size=8)

    def configuration(self) -> None:
        ev = self.ev
        self.h1("Onboard configuration")
        rows = [("Scenario", ev.scenario_title), ("Antenna and modem", ev.vehicle), ("Link management", ev.policy),
                ("Mobile networks", ", ".join(ev.operators)), ("Satellite", ev.satcom), ("Passenger load", ev.passengers)]
        d = ev.design
        if d:
            rows += [("Train design", f"{d.get('title')}: {d.get('n_carriages')} carriages"),
                     ("Roof units", f"{d.get('cellular_units')} EDGE Rail cellular unit(s); {d.get('satcom_units')} satellite terminal(s)"
                                    + (f" ({d.get('satcom_terminal')})" if d.get("satcom_terminal") else "")),
                     ("Passenger Wi-Fi", f"{d.get('aps_connected')} of {d.get('aps_total')} access points connected; Fleet Connect "
                                         f"{'fitted' if d.get('fleet_connect') else 'not fitted'}; {d.get('passengers')} seats")]
        self.kv(rows, "Onboard system modelled")
        if d and d.get("warnings"):
            self.para("Notes on the train design:", keep=True)
            self.bullets(list(d["warnings"]))
        if ev.claims:
            self.para("Manufacturer performance claims for this equipment (" + "; ".join(ev.claims) + ") are shown for reference only; they do "
                      "not enter the model, which uses the link-budget and throughput assumptions in section 9.", size=9, color=MUTED)

    def methodology(self) -> None:
        self.h1("Methodology")
        for title, text in [
            ("Route and journey", "The railway centreline, stations, tunnels and cuttings come from OpenStreetMap, routed station to station. The "
                                  "route is sampled every 50 m; a speed model driven by line speed, station stops and the timetable gives the "
                                  "time the train passes each point."),
            ("Mobile networks", "For each operator at each point, the model starts from the operator's predicted coverage (Ofcom), then applies "
                                "losses for deep cuttings and for distance from the serving cell site, handover effects as the train passes "
                                "from one site to the next, and the gain or loss of the train's antenna. Tunnels have no coverage unless an "
                                "in-tunnel system is listed. Signal (RSRP), throughput, latency and packet loss follow from the resulting quality."),
            ("Satellite", "Sky visibility is calculated from a 30 m terrain model's horizon around each point, reduced for cuttings, station "
                          "canopies and urban obstruction; tunnels block the link. Availability, throughput and latency follow from sky visibility "
                          "and the terminal's characteristics."),
            ("Onboard link management", "The onboard router scores every link on capacity, latency, loss, stability and confidence, and "
                                        "combines the usable ones according to the policy: packet bonding aggregates them, failover uses the best "
                                        "single link."),
            ("Passenger experience", "The combined throughput is shared among the active passengers (seats × load factor × share using Wi-Fi), "
                                     "capped by the access points' capacity. A score from per-passenger throughput, latency and loss gives the "
                                     "service class."),
            ("Confidence", "Each estimate carries a confidence from 0 to 1 reflecting its inputs: synthetic stand-ins score low, published "
                           "predictions moderate, and calibrated or measured values high (section 7)."),
        ]:
            self.h2(title)
            self.para(text)

    def provenance(self) -> None:
        ev = self.ev
        self.h1("Data sources and confidence")
        status = {"live": "Live", "synthetic stand-in": "Stand-in (synthetic)", "predictive": "Predictive model", "configured": "Configured"}
        self.table([("Input", 5.4, "l"), ("Source", 7.6, "l"), ("Status", 3.6, "l")],
                   [[a, b, status.get(c, c)] for a, b, c in ev.sources], "Data sources behind this assessment", size=8.5)
        stand_ins = [a for a, _, c in ev.sources if c == "synthetic stand-in"]
        if stand_ins:
            self.para("Stand-in data was used for: " + "; ".join(stand_ins).lower() + ". Estimates that rest on it carry a lower confidence.",
                      color=MUTED, size=9)
        self.para("Ofcom coverage is operator-predicted, not measured. OpenCellID is community-contributed, so a missing site does not mean "
                  "there is no service. There is no public route-level satellite telemetry, so the satellite model is predictive until "
                  "terminal logs are attached.", keep=True)
        self.table([("Confidence", 3.0, "l"), ("Meaning", 13.6, "l")], [list(r) for r in CONFIDENCE_SCALE], "How to read confidence values", bold_first=True)

    def validation(self) -> None:
        ev = self.ev
        self.h1("Validation status")
        v = ev.validation
        if v is not None and len(v):
            self.para("Predicted values were compared with field measurements attached to the route (matched within 250 m). The table gives "
                      "the agreement by network and route section; the Excel appendix holds every row.", keep=True)
            head = v.head(40)
            self.table([("Network", 2.2, "l"), ("From km", 1.6, "r"), ("Points", 1.5, "r"), ("RSRP MAE dB", 2.0, "r"), ("RSRP RMSE dB", 2.1, "r"),
                        ("Class accuracy", 2.1, "r"), ("Outage precision", 2.2, "r"), ("Outage recall", 2.0, "r"), ("Availability error pp", 2.4, "r")],
                       [[ev.links.set_index("link")["name"].get(r["provider_id"], r["provider_id"]) if "provider_id" in r else "", _fmt(r.get("section_from_km"), ".0f"),
                         _fmt(r.get("n"), ".0f"), _fmt(r.get("rsrp_mae"), ".1f"), _fmt(r.get("rsrp_rmse"), ".1f"),
                         _fmt(100 * r["class_acc"], ".0f") + " %" if pd.notna(r.get("class_acc")) else "–", _fmt(r.get("outage_precision"), ".2f"),
                         _fmt(r.get("outage_recall"), ".2f"), _fmt(r.get("avail_err_pp"), ".1f")] for _, r in head.iterrows()],
                       "Agreement between predictions and field measurements" + (f" (first 40 of {len(v)} rows)" if len(v) > 40 else ""), size=8)
        else:
            self.para("No field measurements have yet been attached to this route, so the figures in this document are unvalidated "
                      "predictions. The calibration and validation tools accept Ofcom drive-test data, the Ofcom Connectivity on Trains study, "
                      "network survey logs and onboard modem logs. Once supplied, this section reports signal error, outage detection, "
                      "service-class accuracy and handover position error by section, and the confidence values rise accordingly.")

    def assumptions(self) -> None:
        ev = self.ev
        self.h1("Assumptions and limitations")
        self.h2("Model assumptions")
        self.table([("Parameter", 5.6, "l"), ("Value", 11.0, "l")], [[a, b] for a, b in ev.assumptions], "Model parameters used for this scenario",
                   size=8, bold_first=True)
        self.h2("Limitations")
        self.bullets([
            "The figures are predictions from published coverage predictions and a simulation; they are not measurements of a deployed system.",
            "Real signal and throughput vary with network load, spectrum, weather and time of day; the model represents these with average assumptions.",
            "Operators' coverage predictions are outdoor predictions; the model adjusts them for the train's antenna and the railway's cuttings and tunnels.",
            "Cell-site data is community-contributed and incomplete, which affects serving-cell distance and handover positions.",
            "Satellite performance is predicted from sky visibility; beam capacity and network load are represented by a fixed prior.",
            "Passenger demand uses a fixed load profile and share of passengers online; peak loads may differ.",
        ])

    def appendix_sections(self) -> None:
        ev = self.ev
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        self._page(sec, landscape=True)
        self.h1("Station-to-station results", number="Appendix A")
        self.para("Full predicted results for each section between consecutive calling points (combined onboard connection).", keep=True)
        self.table([("Section", 6.4, "l"), ("km", 1.3, "r"), ("min", 1.2, "r"), ("Speed km/h", 1.6, "r"), ("Mean Mbps", 1.6, "r"),
                    ("P10 Mbps", 1.5, "r"), ("Min Mbps", 1.5, "r"), ("Per user Mbps", 1.7, "r"), ("Latency ms", 1.6, "r"),
                    ("Streaming %", 1.8, "r"), ("Usable %", 1.6, "r"), ("Loss km", 1.4, "r"), ("Tunnels", 1.4, "r"),
                    ("Least available", 2.1, "l"), ("Conf.", 1.2, "r")],
                   [[f"{r['from']} – {r['to']}", _fmt(r["length_km"], ".1f"), _fmt(r["minutes"], ".0f"), _fmt(r["mean_speed_kph"], ".0f"),
                     _fmt(r["bonded_mean_mbps"], ".0f"), _fmt(r["bonded_p10_mbps"], ".0f"), _fmt(r["bonded_min_mbps"], ".0f"),
                     _fmt(r["per_user_mean_mbps"], ".2f"), _fmt(r["latency_mean_ms"], ".0f"), _fmt(r["streaming_share_pct"], ".1f"),
                     _fmt(r["usable_share_pct"], ".1f"), _fmt(r["outage_km"], ".2f"), _fmt(r["tunnels"]), r["weakest_name"],
                     _fmt(r["confidence"], ".2f")] for _, r in ev.sec.iterrows()],
                   "Predicted results by station-to-station section", size=7.5, total_cm=LANDSCAPE_TEXT_CM)

    def glossary(self) -> None:
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        self._page(sec, landscape=False)
        self.h1("Glossary", number="Appendix B")
        self.table([("Term", 4.4, "l"), ("Meaning", 12.2, "l")], [list(g) for g in GLOSSARY], "Terms used in this document", size=8.5, bold_first=True)

    def build(self, path: Path) -> None:
        first = self.doc.sections[0]
        self._page(first, landscape=False)
        first.different_first_page_header_footer = True     # the cover has no header or footer
        self.cover()
        self.document_control()
        anchor = self.contents_page()
        self.executive_summary()
        self.introduction()
        self.performance()
        self.sections_and_links()
        self.configuration()
        self.methodology()
        self.provenance()
        self.validation()
        self.assumptions()
        self.appendix_sections()
        self.glossary()
        self.contents(anchor)
        self.doc.save(path)


def write_docx(path: Path, ev: Evidence) -> str:
    """Write the report; returns its document reference (the Excel appendix quotes it)."""
    r = _Report(ev)
    r.build(path)
    return r.reference
