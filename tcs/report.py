"""Evidence pack for tender submissions: a Word report (charts, KPIs, station-to-station tables, method, provenance,
validation status) and an Excel data appendix, for one route and one scenario.

Everything in it is a model prediction unless the validation section says otherwise; the wording is deliberate.
"""
from __future__ import annotations

import io
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import pandas as pd

from .config import Settings

WIFI_CLASSES = ["EXCELLENT", "GOOD", "USABLE", "POOR", "OUTAGE"]
CLASS_COLORS = {"EXCELLENT": "#1e7d3e", "GOOD": "#3f9a52", "USABLE": "#c9a100", "POOR": "#d6531a", "OUTAGE": "#9c1b2c"}

# Route heat maps: five bands per metric on one ordinal ramp, darker = worse, so weak stretches stand out. The ramp is
# a single hue with monotone lightness and its light end at 2.1:1 on white (checked with an OKLab validator), so it
# survives greyscale printing and colour-vision deficiency; the legends name every band in words and numbers.
BANDS = ["Excellent", "Good", "Fair", "Poor", "Very poor"]
BAND_COLORS = ["#e8a577", "#db773c", "#c24c0c", "#9d2c04", "#6f1a11"]
HEAT_METRICS = {   # band edges, best first; industry-conventional RSRP bands for LTE/5G
    "signal": {"title": "Signal strength", "sub": "strongest mobile network at the train (RSRP)",
               "ranges": ["≥ −80 dBm", "−90 to −80 dBm", "−100 to −90 dBm", "−110 to −100 dBm", "< −110 dBm or none"]},
    "throughput": {"title": "Throughput", "sub": "combined onboard WAN capacity",
                   "ranges": ["≥ 200 Mbps", "100 to 200 Mbps", "50 to 100 Mbps", "10 to 50 Mbps", "< 10 Mbps or no link"]},
    "latency": {"title": "Latency", "sub": "combined onboard WAN round trip",
                "ranges": ["< 40 ms", "40 to 60 ms", "60 to 100 ms", "100 to 150 ms", "≥ 150 ms or no link"]},
}
INK, MUTED, HAIRLINE = "#1f2933", "#5d6c7b", "#dde3e9"


# ---------------------------------------------------------------- analysis
def section_table(samples: pd.DataFrame, rc: pd.DataFrame, obs: pd.DataFrame, stations: pd.DataFrame, providers: list[str]) -> pd.DataFrame:
    """Station-to-station sections (between consecutive stations in the sequence)."""
    st = stations.sort_values("distance_m").reset_index(drop=True)
    rc = rc.set_index("sample_id")
    rows = []
    wide = obs.pivot_table(index="sample_id", columns="provider_id", values="capacity_mbps")
    avail = obs.pivot_table(index="sample_id", columns="provider_id", values="available", aggfunc="max")
    for k in range(len(st) - 1):
        a, b = int(st.loc[k, "sample_id"]), int(st.loc[k + 1, "sample_id"])
        if b <= a:
            continue
        seg = rc.loc[a:b]
        s = samples.iloc[a:b + 1]
        cls = seg["service_class"].value_counts(normalize=True)
        worst = avail.loc[a:b].mean().idxmin() if len(avail) else ""
        rows.append({
            "section": f"{st.loc[k, 'crs']} → {st.loc[k + 1, 'crs']}", "from": st.loc[k, "name"], "to": st.loc[k + 1, "name"],
            "length_km": round((s["distance_m"].iloc[-1] - s["distance_m"].iloc[0]) / 1000, 1),
            "minutes": round((s["sim_seconds"].iloc[-1] - s["sim_seconds"].iloc[0]) / 60, 1),
            "mean_speed_kph": round(float(s["speed_kph"].mean()), 0),
            "bonded_mean_mbps": round(float(seg["bonded_capacity_mbps"].mean()), 0),
            "bonded_p10_mbps": round(float(seg["bonded_capacity_mbps"].quantile(0.1)), 0),
            "bonded_min_mbps": round(float(seg["bonded_capacity_mbps"].min()), 0),
            "per_user_mean_mbps": round(float(seg["per_user_mbps"].mean()), 2),
            "latency_mean_ms": round(float(seg["effective_latency_ms"].mean()), 0),
            "streaming_share_pct": round(100 * (cls.get("EXCELLENT", 0) + cls.get("GOOD", 0)), 1),
            "usable_share_pct": round(100 * (1 - cls.get("OUTAGE", 0) - cls.get("POOR", 0)), 1),
            "outage_km": round(float(np.sum(np.diff(s["distance_m"]) * (seg["service_class"].iloc[:-1].to_numpy() == "OUTAGE"))) / 1000, 2),
            "tunnels": int(((s["in_tunnel"].astype(int).diff() == 1).sum()) + (1 if s["in_tunnel"].iloc[0] else 0)),
            "weakest_link": str(worst),
            **{f"{p}_avail_pct": round(100 * float(avail.loc[a:b, p].mean()), 1) for p in providers if p in avail},
            **{f"{p}_p50_mbps": round(float(wide.loc[a:b, p].median()), 0) for p in providers if p in wide},
            "confidence": round(float(seg["confidence"].mean()), 2),
        })
    return pd.DataFrame(rows)


def link_table(obs: pd.DataFrame, spacing_m: float) -> pd.DataFrame:
    rows = []
    for pid, g in obs.groupby("provider_id"):
        av = g[g["available"]]
        rows.append({"link": pid, "type": g["provider_type"].iloc[0], "availability_pct": round(100 * g["available"].mean(), 1),
                     "p50_capacity_mbps": round(float(av["capacity_mbps"].median()), 0) if len(av) else 0, "p10_capacity_mbps": round(float(av["capacity_mbps"].quantile(0.1)), 0) if len(av) else 0,
                     "mean_latency_ms": round(float(av["latency_ms"].mean()), 0) if len(av) else np.nan, "handover_events": int(g["handover"].sum()),
                     "unavailable_km": round(float((~g["available"]).sum() * spacing_m / 1000), 1), "mean_confidence": round(float(g["confidence"].mean()), 2),
                     "sources": " | ".join(sorted({f for x in g["source_flags"].astype(str).head(2000) for f in x.split("|")}))})
    return pd.DataFrame(rows)


def kpis(rc: pd.DataFrame, samples: pd.DataFrame, spacing_m: float) -> dict:
    cls = rc["service_class"].value_counts(normalize=True)
    return {
        "route_length_km": round(float(samples["distance_m"].max()) / 1000, 1),
        "journey_minutes": round(float(samples["sim_seconds"].max()) / 60, 0),
        "streaming_share_pct": round(100 * (cls.get("EXCELLENT", 0) + cls.get("GOOD", 0)), 1),
        "usable_share_pct": round(100 * (1 - cls.get("OUTAGE", 0) - cls.get("POOR", 0)), 1),
        "outage_share_pct": round(100 * cls.get("OUTAGE", 0), 2),
        "outage_km": round(float(np.sum(np.diff(samples["distance_m"]) * (rc["service_class"].iloc[:-1].to_numpy() == "OUTAGE"))) / 1000, 1),
        "bonded_mean_mbps": round(float(rc["bonded_capacity_mbps"].mean()), 0),
        "bonded_median_mbps": round(float(rc["bonded_capacity_mbps"].median()), 0),
        "bonded_p10_mbps": round(float(rc["bonded_capacity_mbps"].quantile(0.1)), 0),
        "per_user_mean_mbps": round(float(rc["per_user_mbps"].mean()), 2),
        "latency_mean_ms": round(float(rc["effective_latency_ms"].mean()), 0),
        "wifi_score_mean": round(float(rc["wifi_service_score"].mean()), 0),
        "mean_confidence": round(float(rc["confidence"].mean()), 2),
        "class_share": {c: round(100 * float(cls.get(c, 0)), 1) for c in WIFI_CLASSES},
    }


def rsrp_band(dbm: np.ndarray) -> np.ndarray:
    """0 = Excellent … 4 = Very poor; no signal (NaN) is Very poor."""
    return np.select([dbm >= -80, dbm >= -90, dbm >= -100, dbm >= -110], [0, 1, 2, 3], 4)


def heat_bands(samples: pd.DataFrame, rc: pd.DataFrame, obs: pd.DataFrame) -> dict[str, np.ndarray]:
    """Per-sample band (0 = Excellent … 4 = Very poor) for each HEAT_METRICS entry, in route order."""
    sid = samples["sample_id"].to_numpy()
    cell = obs[obs["provider_type"] == "cellular"]
    best = cell.groupby("sample_id")["signal_primary"].max().reindex(sid).to_numpy(float)
    r = rc.set_index("sample_id").reindex(sid)
    cap, lat = r["bonded_capacity_mbps"].to_numpy(float), r["effective_latency_ms"].to_numpy(float)
    return {"signal": rsrp_band(best),
            "throughput": np.select([cap >= 200, cap >= 100, cap >= 50, cap >= 10], [0, 1, 2, 3], 4),
            "latency": np.select([lat < 40, lat < 60, lat < 100, lat < 150], [0, 1, 2, 3], 4)}   # NaN latency = no link


def typical_bands(bands: np.ndarray, n_runs: int) -> tuple[np.ndarray, np.ndarray]:
    """Split the route into up to n_runs equal runs of samples; each run takes the band that at least half of it reaches
    (the upper median), so charts at page scale show sustained stretches without one 50 m sample colouring a kilometre.
    Returns (run start indices plus the end, band per run)."""
    edges = np.linspace(0, len(bands), min(n_runs, len(bands)) + 1).astype(int)
    typical = np.array([np.sort(bands[a:b])[(b - a) // 2] for a, b in zip(edges[:-1], edges[1:])], dtype=int)
    return edges, typical


def band_shares(bands: np.ndarray) -> list[float]:
    """Share of the route (%) in each band; samples are evenly spaced, so a share of samples is a share of length."""
    return list(np.bincount(bands, minlength=len(BANDS)) / max(len(bands), 1) * 100)


# ---------------------------------------------------------------- charts
def _chart_capacity(samples, rc, stations, title) -> bytes:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    fig, ax = plt.subplots(figsize=(10, 3.2), dpi=160)
    km = samples["distance_m"].values / 1000
    ax.fill_between(km, 0, rc["bonded_capacity_mbps"].values, color="#2457d6", alpha=0.18, linewidth=0)
    ax.plot(km, rc["bonded_capacity_mbps"].values, color="#2457d6", linewidth=0.8)
    for _, s in stations.iterrows():
        ax.axvline(s["distance_m"] / 1000, color="#b9c4cf", linewidth=0.6)
        ax.text(s["distance_m"] / 1000, ax.get_ylim()[1] * 0.98 if ax.get_ylim()[1] > 0 else 1, s["crs"], rotation=90, va="top", ha="right", fontsize=6, color="#5d6c7b")
    tun = samples["in_tunnel"].values.astype(bool)
    if tun.any():
        ax.fill_between(km, 0, ax.get_ylim()[1], where=tun, color="#9c1b2c", alpha=0.25, linewidth=0, label="tunnel")
    import textwrap
    ax.set_xlabel("Distance from origin (km)"); ax.set_ylabel("Combined WAN capacity (Mbps)"); ax.set_title(textwrap.shorten(title, 110, placeholder="…"), fontsize=9, loc="left")
    ax.grid(alpha=0.25); ax.spines[["top", "right"]].set_visible(False)
    buf = io.BytesIO(); fig.tight_layout(); fig.savefig(buf, format="png"); plt.close(fig)
    return buf.getvalue()


def _chart_sections(sec: pd.DataFrame) -> bytes:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    fig, ax = plt.subplots(figsize=(10, 3.0), dpi=160)
    x = np.arange(len(sec))
    ax.bar(x, sec["streaming_share_pct"], color=CLASS_COLORS["GOOD"], label="video-call / streaming capable")
    ax.bar(x, sec["usable_share_pct"] - sec["streaming_share_pct"], bottom=sec["streaming_share_pct"], color=CLASS_COLORS["USABLE"], label="usable")
    ax.bar(x, 100 - sec["usable_share_pct"], bottom=sec["usable_share_pct"], color=CLASS_COLORS["OUTAGE"], label="poor / outage")
    ax.set_xticks(x); ax.set_xticklabels(sec["section"], rotation=60, ha="right", fontsize=6)
    ax.set_ylabel("% of section length"); ax.set_ylim(0, 100); ax.legend(fontsize=7, frameon=False, loc="lower left")
    ax.set_title("Predicted passenger Wi-Fi service class by station-to-station section", fontsize=10, loc="left")
    ax.spines[["top", "right"]].set_visible(False)
    buf = io.BytesIO(); fig.tight_layout(); fig.savefig(buf, format="png"); plt.close(fig)
    return buf.getvalue()


def _chart_links(links: pd.DataFrame) -> bytes:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    fig, axes = plt.subplots(1, 2, figsize=(10, 2.8), dpi=160)
    axes[0].barh(links["link"], links["availability_pct"], color="#2457d6"); axes[0].set_xlim(0, 100); axes[0].set_title("Availability (% of route)", fontsize=9, loc="left")
    axes[1].barh(links["link"], links["p50_capacity_mbps"], color="#6f9bd8"); axes[1].set_title("Median capacity when available (Mbps)", fontsize=9, loc="left")
    for ax in axes:
        ax.spines[["top", "right"]].set_visible(False); ax.grid(axis="x", alpha=0.25)
    buf = io.BytesIO(); fig.tight_layout(); fig.savefig(buf, format="png"); plt.close(fig)
    return buf.getvalue()


def _band_legend(ax, metric: str, bands: np.ndarray, ncol: int = 1) -> None:
    from matplotlib.patches import Patch

    shares = band_shares(bands)
    fmt = lambda v: "0 %" if v == 0 else ("<1 %" if v < 0.5 else f"{v:.0f} %")
    labels = [f"{b}  {r}  ·  {fmt(v)}" for b, r, v in zip(BANDS, HEAT_METRICS[metric]["ranges"], shares)]
    handles = [Patch(facecolor=c, edgecolor="none") for c in BAND_COLORS]
    ax.legend(handles, labels, loc="upper left", ncol=ncol, frameon=False, fontsize=6.5, handlelength=1.4, handleheight=0.9,
              borderaxespad=0, labelcolor=INK, title="Share of route length", title_fontsize=6.5, alignment="left")


def _scale_bar(ax) -> None:
    """A km scale bar in the lower-left corner (axes are in km, equal aspect)."""
    x0, x1 = ax.get_xlim(); y0, y1 = ax.get_ylim()
    target = (x1 - x0) * 0.22
    length = max([v for v in (1, 2, 5, 10, 20, 25, 50, 100, 200) if v <= target] or [1])
    bx, by = x0 + (x1 - x0) * 0.04, y0 + (y1 - y0) * 0.05
    ax.plot([bx, bx + length], [by, by], color=MUTED, linewidth=1.2, solid_capstyle="butt", zorder=4)
    ax.text(bx + length / 2, by + (y1 - y0) * 0.02, f"{length} km", ha="center", va="bottom", fontsize=6, color=MUTED, zorder=4)


def _chart_heatmap(samples: pd.DataFrame, stations: pd.DataFrame, bands: dict[str, np.ndarray]) -> bytes:
    """Three small-multiple maps of the railway (signal, throughput, latency), each coloured by band along the line."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.collections import LineCollection
    from matplotlib.patheffects import withStroke

    lat0 = np.radians(float(samples["latitude"].mean()))
    kx, ky = 111.32 * np.cos(lat0), 110.57                    # equirectangular km: accurate to <1 % over a UK route
    x, y = samples["longitude"].to_numpy() * kx, samples["latitude"].to_numpy() * ky
    w, h = max(float(np.ptp(x)), 1.0), max(float(np.ptp(y)), 1.0)
    xy = np.column_stack([x, y])
    tall = h > 1.1 * w                                         # north-south route: maps side by side; otherwise stacked
    if tall:
        map_h = float(np.clip(3.0 * h / w, 2.6, 6.0))
        fig = plt.figure(figsize=(10, map_h + 1.9), dpi=160)
        gs = fig.add_gridspec(2, 3, height_ratios=[map_h, 1.05], left=0.02, right=0.98, top=1 - 0.5 / (map_h + 1.9), bottom=0.01, wspace=0.08, hspace=0.06)
        cells = [(gs[0, i], gs[1, i]) for i in range(3)]
    else:
        map_h = float(np.clip(6.4 * h / w, 1.5, 2.6))
        fig = plt.figure(figsize=(10, 3 * (map_h + 0.5)), dpi=160)
        gs = fig.add_gridspec(3, 2, width_ratios=[6.4, 3.0], left=0.02, right=0.98, top=1 - 0.45 / (3 * (map_h + 0.5)), bottom=0.01, hspace=0.32, wspace=0.04)
        cells = [(gs[i, 0], gs[i, 1]) for i in range(3)]
    # Page scale: a run of line is about 2 pt long, and neighbouring runs overlap by about 0.7 pt so anti-aliasing leaves no seams.
    box_w, box_h = (3.0, map_h) if tall else (6.3, map_h)      # map box in inches (approximate; the 6 % padding is below)
    pt_per_km = 72 * min(box_w / (w * 1.12), box_h / (h * 1.12))
    path_km = float(np.hypot(np.diff(x), np.diff(y)).sum())
    n_runs = int(np.clip(path_km * pt_per_km / 2.0, 50, 600))
    spacing_km = max(float(np.median(np.diff(samples["distance_m"].to_numpy()))) / 1000, 1e-3)
    overlap = max(1, int(np.ceil(0.7 / pt_per_km / spacing_km)))
    stops = stations[stations["stop"]] if "stop" in stations else stations
    sx, sy = stops["longitude"].to_numpy() * kx, stops["latitude"].to_numpy() * ky
    for (map_cell, legend_cell), (metric, spec) in zip(cells, HEAT_METRICS.items()):
        ax = fig.add_subplot(map_cell)
        ax.set_aspect("equal", adjustable="datalim")
        pad = 0.06 * max(w, h)
        ax.set_xlim(x.min() - pad, x.max() + pad); ax.set_ylim(y.min() - pad, y.max() + pad)
        ax.plot(x, y, color="#8b939c", linewidth=4.4, solid_capstyle="round", solid_joinstyle="round", zorder=1)   # casing: the palest band still reads
        edges, typical = typical_bands(bands[metric], n_runs)
        change = np.r_[True, typical[1:] != typical[:-1]]      # merge neighbouring runs in the same band
        starts, kinds = edges[:-1][change], typical[change]
        ends = np.r_[starts[1:], len(xy)]
        runs = [xy[max(a - overlap, 0):min(b + overlap, len(xy))] for a, b in zip(starts, ends)]
        order = np.argsort(kinds, kind="stable")               # worse runs on top where runs overlap
        ax.add_collection(LineCollection([runs[i] for i in order], colors=np.array(BAND_COLORS)[kinds[order]], linewidths=3.0,
                                         capstyle="butt", joinstyle="round", zorder=2))
        ax.scatter(sx, sy, s=14, facecolor="white", edgecolor=INK, linewidth=0.8, zorder=3)
        placed: list[tuple[float, float]] = []
        gap = 0.07 * max(w, h)
        for k in [0, len(stops) - 1, *range(1, len(stops) - 1)]:   # ends first, then intermediate stops that have room
            if k < 0 or any(np.hypot(sx[k] - px, sy[k] - py) < gap for px, py in placed):
                continue
            placed.append((sx[k], sy[k]))
            ax.annotate(stops["crs"].iloc[k], (sx[k], sy[k]), xytext=(4, 3), textcoords="offset points", fontsize=6, color=INK, zorder=4,
                        path_effects=[withStroke(linewidth=2.2, foreground="white")])
        ax.set_xticks([]); ax.set_yticks([])
        for sp in ax.spines.values():
            sp.set_color(HAIRLINE); sp.set_linewidth(0.6)
        ax.annotate("N", xy=(0.96, 0.96), xytext=(0, -16), xycoords="axes fraction", textcoords="offset points", ha="center", va="center",
                    fontsize=6.5, color=MUTED, arrowprops={"arrowstyle": "-|>", "color": MUTED, "lw": 0.8})
        ax.annotate(spec["title"], (0, 1), xytext=(0, 12), xycoords="axes fraction", textcoords="offset points", fontsize=9, fontweight="bold", color=INK, va="bottom")
        ax.annotate(spec["sub"], (0, 1), xytext=(0, 3), xycoords="axes fraction", textcoords="offset points", fontsize=6.5, color=MUTED, va="bottom")
        fig.canvas.draw()                                      # settle the equal-aspect limits before placing the scale bar
        _scale_bar(ax)
        lax = fig.add_subplot(legend_cell); lax.axis("off")
        _band_legend(lax, metric, bands[metric])
    buf = io.BytesIO(); fig.savefig(buf, format="png"); plt.close(fig)
    return buf.getvalue()


def _chart_networks(samples: pd.DataFrame, stations: pd.DataFrame, obs: pd.DataFrame, names: dict[str, str], best: np.ndarray) -> bytes:
    """Heat strip: signal band of each mobile network (rows) along the route (columns), plus the strongest of them."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.colors import ListedColormap
    from matplotlib.patches import Patch

    sid = samples["sample_id"].to_numpy()
    cell = obs[obs["provider_type"] == "cellular"]
    ids = list(dict.fromkeys(cell["provider_id"]))
    rows = [rsrp_band(cell[cell["provider_id"] == p].set_index("sample_id")["signal_primary"].reindex(sid).to_numpy(float)) for p in ids]
    rows.append(best)
    labels = [names.get(p, p) for p in ids] + ["Strongest"]
    grid = np.vstack([typical_bands(r, 500)[1] for r in rows])  # same smoothing as the maps: about 0.3 mm of print per column
    km = float(samples["distance_m"].max()) / 1000
    fig, ax = plt.subplots(figsize=(10, 0.34 * len(rows) + 1.35), dpi=160)
    ax.imshow(grid, cmap=ListedColormap(BAND_COLORS), vmin=-0.5, vmax=len(BANDS) - 0.5, aspect="auto", interpolation="nearest", extent=[0, km, len(rows), 0])
    for i in range(1, len(rows)):
        ax.axhline(i, color="white", linewidth=2.2 if i == len(rows) - 1 else 1.2)   # a wider gap sets off the "Strongest" row
    ax.set_yticks(np.arange(len(rows)) + 0.5); ax.set_yticklabels(labels, fontsize=7, color=INK)
    ax.tick_params(axis="y", length=0)
    ax.set_xlabel("Distance from origin (km)", fontsize=7, color=MUTED); ax.tick_params(axis="x", labelsize=6.5, colors=MUTED)
    stops = stations[stations["stop"]] if "stop" in stations else stations
    ax.set_xlim(0, km)
    top = ax.secondary_xaxis("top")
    top.set_xticks(np.minimum(stops["distance_m"].to_numpy() / 1000, km)); top.set_xticklabels(stops["crs"], fontsize=5.5, rotation=90, color=MUTED)
    top.tick_params(length=2, colors=MUTED)
    for sp in ax.spines.values():
        sp.set_visible(False)
    for sp in top.spines.values():
        sp.set_visible(False)
    ax.legend([Patch(facecolor=c) for c in BAND_COLORS], [f"{b}  {r}" for b, r in zip(BANDS, HEAT_METRICS["signal"]["ranges"])], loc="upper center",
              bbox_to_anchor=(0.5, -0.32 if len(rows) > 3 else -0.5), ncol=5, frameon=False, fontsize=6.5, handlelength=1.4, labelcolor=INK)
    buf = io.BytesIO(); fig.tight_layout(); fig.savefig(buf, format="png", bbox_inches="tight", pad_inches=0.08); plt.close(fig)
    return buf.getvalue()


# ---------------------------------------------------------------- documents
def write_xlsx(path: Path, k: dict, sec: pd.DataFrame, links: pd.DataFrame, samples: pd.DataFrame, rc: pd.DataFrame, obs: pd.DataFrame, assumptions: list[tuple[str, str]], sources: list[tuple[str, str, str]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head = Font(bold=True, color="FFFFFF"); fill = PatternFill("solid", fgColor="2457D6")

    def sheet(name, df: pd.DataFrame):
        ws = wb.create_sheet(name)
        ws.append(list(df.columns))
        for c in ws[1]:
            c.font = head; c.fill = fill; c.alignment = Alignment(vertical="center")
        for row in df.itertuples(index=False):
            ws.append([None if (isinstance(v, float) and np.isnan(v)) else (v.item() if hasattr(v, "item") else v) for v in row])
        for i, col in enumerate(df.columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = min(48, max(10, len(str(col)) + 2))
        ws.freeze_panes = "A2"
        return ws

    ws = wb.active; ws.title = "Summary"
    ws.append(["Metric", "Value"]); ws["A1"].font = head; ws["A1"].fill = fill; ws["B1"].font = head; ws["B1"].fill = fill
    for kk, v in k.items():
        if kk == "class_share":
            for c, share in v.items():
                ws.append([f"share_{c.lower()}_pct", share])
        else:
            ws.append([kk, v])
    ws.column_dimensions["A"].width = 30; ws.column_dimensions["B"].width = 16
    sheet("Sections", sec)
    sheet("Links", links)
    wide = obs.pivot_table(index="sample_id", columns="provider_id", values=["capacity_mbps", "latency_ms", "available", "quality_score"])
    wide.columns = [f"{p}_{m}" for m, p in wide.columns]
    smp = samples[["sample_id", "distance_m", "timestamp_sim", "latitude", "longitude", "speed_kph", "elevation_m", "in_tunnel", "cutting_depth_m", "sky_visibility", "station_nearby"]].copy()
    smp["timestamp_sim"] = pd.to_datetime(smp["timestamp_sim"]).dt.strftime("%H:%M:%S")
    smp = smp.merge(wide.reset_index(), on="sample_id", how="left").merge(
        rc[["sample_id", "active_links", "bonded_capacity_mbps", "effective_latency_ms", "packet_loss_pct", "per_user_mbps", "wifi_service_score", "service_class", "confidence", "source_flags"]], on="sample_id", how="left")
    sheet("Samples", smp)
    sheet("Assumptions", pd.DataFrame(assumptions, columns=["parameter", "value"]))
    sheet("Sources", pd.DataFrame(sources, columns=["layer", "source", "status"]))
    wb.save(path)


def write_docx(path: Path, settings: Settings, meta: dict, k: dict, sec: pd.DataFrame, links: pd.DataFrame, charts: dict[str, bytes], assumptions: list[tuple[str, str]],
               sources: list[tuple[str, str, str]], validation: pd.DataFrame | None, scenario_label: str, design: dict | None) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2); s.top_margin = s.bottom_margin = Cm(1.8)
    style = doc.styles["Normal"]; style.font.name = "Calibri"; style.font.size = Pt(10)

    def table(df: pd.DataFrame, cols: list[str] | None = None, headers: dict[str, str] | None = None, font=8):
        cols = cols or list(df.columns)
        t = doc.add_table(rows=1, cols=len(cols)); t.style = "Light Grid Accent 1"
        for i, c in enumerate(cols):
            cell = t.rows[0].cells[i]; cell.text = (headers or {}).get(c, c); cell.paragraphs[0].runs[0].font.size = Pt(font); cell.paragraphs[0].runs[0].font.bold = True
        for _, r in df.iterrows():
            cells = t.add_row().cells
            for i, c in enumerate(cols):
                v = r[c]; cells[i].text = "" if (isinstance(v, float) and np.isnan(v)) else (f"{v:,.0f}" if isinstance(v, (float, np.floating)) and abs(v) >= 100 else str(v))
                cells[i].paragraphs[0].runs[0].font.size = Pt(font)
        return t

    route = meta["route"]
    doc.add_heading("Onboard connectivity performance evidence", 0)
    doc.add_paragraph(f"{route['name']} · {meta['stations'][0]['name']} → {meta['stations'][-1]['name']} · {k['route_length_km']} km · {int(k['journey_minutes'])} min")
    p = doc.add_paragraph(); r = p.add_run(f"Scenario: {scenario_label}"); r.bold = True
    doc.add_paragraph(f"Generated {datetime.now(timezone.utc).strftime('%d %B %Y %H:%M UTC')} · model version {meta['model_version']} · sample spacing {route['sample_spacing_m']} m ({meta['n_samples']:,} samples)")
    box = doc.add_paragraph()
    rr = box.add_run("Status of this evidence: model prediction. ")
    rr.bold = True; rr.font.color.rgb = RGBColor(0x9C, 0x1B, 0x2C)
    box.add_run("Figures below are simulation outputs built from the route geometry, terrain, published coverage predictions and the onboard architecture described in the method section. "
                "They are not measurements of the deployed system. Confidence values state how much of each figure rests on measured, predicted or synthetic inputs; the validation section states what has been checked against field data.")

    num = iter(range(1, 20))                                   # section numbers (the architecture section is optional)
    doc.add_heading(f"{next(num)}. Headline results", 1)
    kp = pd.DataFrame([
        ("Share of route supporting video calls / streaming (EXCELLENT or GOOD)", f"{k['streaming_share_pct']} %"),
        ("Share of route usable or better", f"{k['usable_share_pct']} %"),
        ("Predicted outage (no usable WAN link)", f"{k['outage_km']} km ({k['outage_share_pct']} %)"),
        ("Combined WAN capacity — mean / median / 10th percentile", f"{k['bonded_mean_mbps']:.0f} / {k['bonded_median_mbps']:.0f} / {k['bonded_p10_mbps']:.0f} Mbps"),
        ("Per-active-user throughput (mean)", f"{k['per_user_mean_mbps']} Mbps"),
        ("Effective latency (mean, when connected)", f"{k['latency_mean_ms']:.0f} ms"),
        ("Passenger Wi-Fi service score (mean, 0–100)", f"{k['wifi_score_mean']:.0f}"),
        ("Mean confidence of the estimate (0–1)", f"{k['mean_confidence']}"),
    ], columns=["Metric", "Value"])
    table(kp, font=9)
    doc.add_paragraph()
    doc.add_picture(io.BytesIO(charts["capacity"]), width=Cm(17))
    doc.add_paragraph("Figure 1. Predicted combined onboard WAN capacity along the route; vertical lines mark stations, shaded bands mark tunnels.").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture(io.BytesIO(charts["sections"]), width=Cm(17))
    doc.add_paragraph("Figure 2. Predicted passenger Wi-Fi service class by station-to-station section.").alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading(f"{next(num)}. Route heat maps", 1)
    doc.add_paragraph("Where along the railway connectivity is strong and where it is weak, darker meaning worse. Signal is the strongest mobile "
                      "network received at the train with this scenario's antenna; throughput and latency are for the combined onboard WAN after "
                      "the link-manager policy, which also draws on the satellite link. At page scale each short stretch of line shows the band "
                      f"that at least half of it reaches; the legend shares count every {route['sample_spacing_m']} m sample, and the Excel "
                      "appendix lists them all. Circles mark calling stations.")
    doc.add_picture(io.BytesIO(charts["heatmap"]), width=Cm(17))
    doc.add_paragraph("Figure 3. Route heat maps: predicted signal strength, throughput and latency (darker = worse).").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture(io.BytesIO(charts["networks"]), width=Cm(17))
    doc.add_paragraph("Figure 4. Predicted signal strength of each mobile network along the route (RSRP at the train, darker = weaker); "
                      "the bottom row is the strongest of them, as mapped in Figure 3.").alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading(f"{next(num)}. Station-to-station performance", 1)
    doc.add_paragraph("Each row is the section between consecutive stations in the service's calling pattern. Capacities are the combined onboard WAN after the link-manager policy; the streaming share is the proportion of the section's length in the EXCELLENT or GOOD passenger Wi-Fi classes.")
    table(sec, ["section", "length_km", "minutes", "bonded_mean_mbps", "bonded_p10_mbps", "per_user_mean_mbps", "latency_mean_ms", "streaming_share_pct", "usable_share_pct", "outage_km", "tunnels", "weakest_link", "confidence"],
          {"section": "Section", "length_km": "km", "minutes": "min", "bonded_mean_mbps": "Mean Mbps", "bonded_p10_mbps": "P10 Mbps", "per_user_mean_mbps": "Per user Mbps", "latency_mean_ms": "Latency ms",
           "streaming_share_pct": "Streaming %", "usable_share_pct": "Usable %", "outage_km": "Outage km", "tunnels": "Tunnels", "weakest_link": "Weakest link", "confidence": "Conf."}, font=7)

    doc.add_heading(f"{next(num)}. Individual links", 1)
    doc.add_picture(io.BytesIO(charts["links"]), width=Cm(17))
    table(links, ["link", "type", "availability_pct", "p50_capacity_mbps", "p10_capacity_mbps", "mean_latency_ms", "handover_events", "unavailable_km", "mean_confidence"],
          {"link": "Link", "type": "Type", "availability_pct": "Available %", "p50_capacity_mbps": "P50 Mbps", "p10_capacity_mbps": "P10 Mbps", "mean_latency_ms": "Latency ms", "handover_events": "Handovers", "unavailable_km": "Unavailable km", "mean_confidence": "Conf."})

    if design:
        doc.add_heading(f"{next(num)}. Onboard architecture", 1)
        doc.add_paragraph(f"Train design: {design.get('title')} — {design.get('n_carriages')} carriages, {design.get('cellular_units')} EDGE Rail cellular roof unit(s), "
                          f"{design.get('satcom_units')} satcom terminal(s){' (' + str(design.get('satcom_terminal')) + ')' if design.get('satcom_terminal') else ''}, "
                          f"{design.get('aps_connected')}/{design.get('aps_total')} access points connected, Fleet Connect {'present' if design.get('fleet_connect') else 'absent'}; "
                          f"{design.get('passengers')} seats modelled. Link policy {str(design.get('policy')).lower().replace('_', ' ')}; vehicle profile {str(design.get('vehicle_profile')).lower().replace('_', ' ')}.")
        for w in design.get("warnings", []):
            doc.add_paragraph(w, style="List Bullet")

    doc.add_heading(f"{next(num)}. Method and assumptions", 1)
    doc.add_paragraph("The route is sampled every 50 m along the railway centreline. At each sample the model estimates every cellular operator (coverage prior → terrain and cutting adjustment → serving-cell distance and handover effects → vehicle/antenna profile) and the satellite link (sky visibility from a digital elevation model horizon, tunnels and station canopies → availability → capacity), then runs the onboard link manager and a passenger demand model to predict the Wi-Fi experience. Every numerical assumption is listed below and versioned in the configuration files that accompany this pack.")
    table(pd.DataFrame(assumptions, columns=["Parameter", "Value"]), font=8)

    doc.add_heading(f"{next(num)}. Data sources and provenance", 1)
    table(pd.DataFrame(sources, columns=["Layer", "Source", "Status"]), font=8)
    doc.add_paragraph("Ofcom coverage is operator-predicted, not measured. OpenCellID is community-contributed; a missing cell does not imply no service. There is no public route-level Starlink RF telemetry; the satellite model is predictive until terminal telemetry is ingested. Throughput depends on network load and spectrum as well as signal.")

    doc.add_heading(f"{next(num)}. Validation status", 1)
    if validation is not None and len(validation):
        doc.add_paragraph("Predicted values were compared with field measurements attached to the route (within 250 m). Metrics by route section:")
        table(validation.head(60), font=7)
    else:
        doc.add_paragraph("No field measurements have yet been attached to this route. The calibration and validation tooling (tcs calibrate / tcs validate) accepts Ofcom drive-test data, the Ofcom Connectivity on Trains study annexes, Network Survey logs and onboard modem logs; once supplied, this section reports MAE/RMSE for signal, outage precision/recall, classification accuracy and handover position error by section, and the confidence values above rise accordingly.")

    doc.add_heading(f"{next(num)}. How to read confidence", 1)
    for line in ["0.90–1.00 directly measured / well validated", "0.70–0.89 strong source coverage + calibrated model", "0.40–0.69 prediction with partial infrastructure support", "0.10–0.39 sparse data / synthetic estimate", "0.00 unknown"]:
        doc.add_paragraph(line, style="List Bullet")
    doc.save(path)


# ---------------------------------------------------------------- entry point
def build_report(settings: Settings, meta: dict, samples: pd.DataFrame, obs: pd.DataFrame, rc: pd.DataFrame, stations: pd.DataFrame, out_dir: Path, scenario_label: str,
                 validation: pd.DataFrame | None = None) -> dict[str, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    providers = list(obs["provider_id"].unique())
    spacing = float(settings.spacing_m)
    k = kpis(rc, samples, spacing)
    sec = section_table(samples, rc, obs, stations, providers)
    links = link_table(obs, spacing)
    sim = settings.sim
    vp = sim["vehicle"]["profiles"][sim["vehicle"]["profile"]]
    assumptions = [
        ("Vehicle / antenna profile", f"{sim['vehicle']['profile']} (score offset {vp.get('score_offset')}, {vp.get('db_offset')} dB, throughput factor ×{vp.get('capacity_factor', 1)})"),
        ("Link-manager policy", sim["wan"]["policy"]), ("Bonding efficiency × congestion factor", f"{sim['wan']['bonding_efficiency']} × {sim['wan']['congestion_factor']}"),
        ("Minimum link score to participate", str(sim["wan"]["minimum_link_score"])), ("Link score weights", str(sim["wan"]["score_weights"])),
        ("Cellular capacity priors (Mbps at excellent signal)", "; ".join(f"{op['id']}: {op['capacity_prior_mbps']}" for op in settings.operators)),
        ("Roof-unit aggregation factor", str(sim["cellular"].get("units_capacity_factor", 1.0))),
        ("Cutting penalty", f"up to −{sim['cellular']['terrain']['cutting_penalty_max']} score at {sim['cellular']['terrain']['cutting_full_depth_m']} m"),
        ("Serving-cell distance penalty", f"−{sim['cellular']['cell_distance']['penalty_per_km']} per km beyond {sim['cellular']['cell_distance']['free_km']} km"),
        ("Handover", f"{sim['cellular']['handover']['duration_samples']} samples, +{sim['cellular']['handover']['latency_spike_ms']} ms, capacity ×{sim['cellular']['handover']['capacity_factor']}"),
        ("Tunnels", f"score {sim['cellular']['tunnels']['default_score']} (no infrastructure) / {sim['cellular']['tunnels']['das_score']} where in-tunnel coverage is assumed: {sim['cellular']['tunnels']['das_tunnels']}"),
        ("Satcom", ", ".join(f"{p['id']}: terminal {p['terminal']}, prior {p['capacity_prior_mbps'][p['terminal']]} Mbps, min elevation {p['min_elevation_deg'][p['terminal']]}°, sky threshold {p['availability']['sky_threshold']}" for p in settings.starlink["satcom"]["providers"]) + ("" if sim.get("satcom_enabled", True) else " (disabled by the train design)")),
        ("Passenger demand", f"{sim['passenger_wifi']['passengers']} seats, load {sim['passenger_wifi']['load_factor_range']}, {sim['passenger_wifi']['active_share']} active, {sim['passenger_wifi']['per_user_demand_mbps']} Mbps each, AP capacity {sim['passenger_wifi']['ap_capacity_mbps']} Mbps"),
        ("Service classes (score lower bounds)", str(sim["passenger_wifi"]["classes"])),
        ("Timetable", f"departure {settings.route.get('timetable', {}).get('departure')} · dwell {settings.route.get('timetable', {}).get('dwell_s')} s · calling pattern from the route configuration"),
    ]
    if sim.get("active_preset") == "edge_rail_fleet_connect":
        pr = sim["presets"]["edge_rail_fleet_connect"]
        assumptions.append(("Manufacturer claims (shown, not modelled)", "; ".join(pr.get("claims", []))))
    live = lambda ok: "live" if ok else "synthetic stand-in"
    cov = meta.get("coverage_sources", [])
    sources = [
        ("Route centreline, stations, tunnels, cuttings, line speed", "OpenStreetMap (ODbL), routed station-to-station" if meta["geometry_source"] == "osm" else meta["geometry_source"], live(meta["geometry_source"] in ("osm", "file"))),
        ("Terrain and sky visibility", meta["terrain_source"], live(meta["terrain_source"] != "synthetic_terrain")),
        ("Cellular coverage prior", ", ".join(cov) if cov else "synthetic", live(any(c.startswith("ofcom") for c in cov))),
        ("Cell sites / handovers", meta.get("cell_source", ""), live(meta.get("cell_source") == "opencellid")),
        ("Satellite", "predictive obstruction model" + (" + terminal telemetry" if any("telemetry" in str(x) for x in obs["source_flags"].unique()) else ""), "predictive"),
        ("Timetable", settings.route.get("timetable", {}).get("source", "yaml"), "configured"),
    ]
    bands = heat_bands(samples, rc, obs)
    names = {op["id"]: op.get("name", op["id"]) for op in settings.operators}
    charts = {"capacity": _chart_capacity(samples, rc, stations, f"{meta['route']['name']} — {scenario_label}"), "sections": _chart_sections(sec), "links": _chart_links(links),
              "heatmap": _chart_heatmap(samples, stations, bands), "networks": _chart_networks(samples, stations, obs, names, bands["signal"])}
    design = sim.get("train", {}).get("design")
    stem = f"evidence_{meta['route']['id']}_{(sim.get('active_preset') or 'baseline')}"
    docx_path, xlsx_path = out_dir / f"{stem}.docx", out_dir / f"{stem}.xlsx"
    write_docx(docx_path, settings, meta, k, sec, links, charts, assumptions, sources, validation, scenario_label, design)
    write_xlsx(xlsx_path, k, sec, links, samples, rc, obs, assumptions, sources)
    for name, png in charts.items():
        (out_dir / f"{stem}_{name}.png").write_bytes(png)
    return {"docx": docx_path, "xlsx": xlsx_path}
